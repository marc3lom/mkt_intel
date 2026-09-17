"""Montagem do documento final a partir do template da mesa.

O template `templates/comentario.dotx` decide o layout: papel, margens, fonte,
marcadores, e as duas imagens de topo e rodapé que fazem o papel de cabeçalho.
Nada aqui inventa formatação — o que este módulo faz é substituir o conteúdo dos
parágrafos de exemplo preservando o `pPr` de cada um, de modo que estilo, bullet
e justificação continuem vindo do template.
"""

from __future__ import annotations

import copy
import re
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.shared import Inches

# Largura útil de texto do template: A4 menos as margens laterais de 1,18".
# É a largura das duas imagens que o template já traz, e a que faz o painel e o
# calendário alinharem com elas.
LARGURA_UTIL = Inches(5.906)

# Todos os runs do corpo do template carregam Aptos explicitamente. O padrão do
# documento é Times New Roman, então run inserido sem fonte sai destoando.
FONTE = "Aptos"

TIPO_TEMPLATE = ("application/vnd.openxmlformats-officedocument."
                 "wordprocessingml.template.main+xml")
TIPO_DOCUMENTO = ("application/vnd.openxmlformats-officedocument."
                  "wordprocessingml.document.main+xml")

MARCA_PAINEL = "[Inserir a tabela de fechamento dos mercados]"
MARCA_CALENDARIO = "[Gráfico do dia]"
RE_PARAGRAFO = re.compile(r"^\[Parágrafo \d+\s*[–-]")

# O guia de estilo fixa o comentário entre quatro e cinco marcadores.
MIN_MARCADORES, MAX_MARCADORES = 4, 5

RE_FECHO = re.compile(r"^\s*(atenciosamente|mesa de investimentos|cordialmente|"
                      r"atenciosos cumprimentos)\b", re.IGNORECASE)


# --------------------------------------------------------------------------
# Markdown
# --------------------------------------------------------------------------

@dataclass(frozen=True)
class Trecho:
    """Um pedaço de texto com sua ênfase."""

    texto: str
    italico: bool = False
    negrito: bool = False


# Alternância com o par duplo antes do simples, para que ** não seja lido como
# dois * vazios. O grupo capturado é sempre o conteúdo interno.
RE_ENFASE = re.compile(
    r"\*\*(?P<forte>[^*]+?)\*\*"
    r"|__(?P<forte2>[^_]+?)__"
    r"|\*(?P<enfase>[^*]+?)\*"
    r"|_(?P<enfase2>[^_]+?)_"
)


def trechos_de(linha: str) -> list[Trecho]:
    """Divide uma linha de Markdown em trechos com ênfase.

    Cobre só o vocabulário que o guia de estilo produz: itálico para termo em
    inglês sem equivalente limpo em português, e negrito eventual. Uma biblioteca
    de Markdown resolveria mais, mas passaria por HTML — conversão que não é
    usada aqui e que traria suas próprias regras de escape.
    """
    saida: list[Trecho] = []
    pos = 0
    for m in RE_ENFASE.finditer(linha):
        if m.start() > pos:
            saida.append(Trecho(linha[pos:m.start()]))
        if m.group("forte") or m.group("forte2"):
            saida.append(Trecho(m.group("forte") or m.group("forte2"), negrito=True))
        else:
            saida.append(Trecho(m.group("enfase") or m.group("enfase2"), italico=True))
        pos = m.end()
    if pos < len(linha):
        saida.append(Trecho(linha[pos:]))
    return [t for t in saida if t.texto]


def le_marcadores(caminho: Path) -> tuple[list[list[Trecho]], list[str]]:
    """Extrai os marcadores do Markdown revisado.

    Devolve os marcadores e as linhas ignoradas, para que nada seja descartado em
    silêncio: título, fecho ou prosa solta fora de marcador viram aviso, não
    sumiço.
    """
    marcadores: list[list[str]] = []
    ignoradas: list[str] = []

    for bruta in caminho.read_text(encoding="utf-8").splitlines():
        linha = bruta.rstrip()
        if not linha.strip():
            continue
        if re.match(r"^\s*[-*+]\s+", linha):
            marcadores.append([re.sub(r"^\s*[-*+]\s+", "", linha)])
        elif marcadores and not linha.startswith("#") and not RE_FECHO.match(linha):
            # Continuação do marcador anterior, quebrada por largura de coluna.
            marcadores[-1].append(linha.strip())
        else:
            ignoradas.append(linha.strip())

    return [trechos_de(" ".join(partes)) for partes in marcadores], ignoradas


# --------------------------------------------------------------------------
# Template
# --------------------------------------------------------------------------

def abre_template(dotx: Path, destino: Path) -> Document:
    """Abre o .dotx como documento editável.

    O python-docx recusa .dotx pelo content type. A conversão é só a troca dessa
    declaração no [Content_Types].xml; o resto do pacote — estilos, numeração,
    imagens, seção — segue idêntico, que é justamente o que se quer preservar.
    """
    with zipfile.ZipFile(dotx) as origem, zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as saida:
        for item in origem.infolist():
            dados = origem.read(item.filename)
            if item.filename == "[Content_Types].xml":
                dados = dados.replace(TIPO_TEMPLATE.encode(), TIPO_DOCUMENTO.encode())
            saida.writestr(item, dados)
    return Document(str(destino))


def _limpa(paragrafo) -> None:
    """Remove o conteúdo do parágrafo, preservando o pPr.

    É o pPr que carrega estilo, bullet e justificação vindos do template. Criar
    um parágrafo novo obrigaria a recriar tudo isso à mão.
    """
    for filho in list(paragrafo._element):
        if not filho.tag.endswith("}pPr"):
            paragrafo._element.remove(filho)


def _escreve(paragrafo, trechos: list[Trecho]) -> None:
    _limpa(paragrafo)
    for t in trechos:
        run = paragrafo.add_run(t.texto)
        run.font.name = FONTE
        run.italic = t.italico or None
        run.bold = t.negrito or None


def _imagem(paragrafo, caminho: Path) -> None:
    _limpa(paragrafo)
    paragrafo.add_run().add_picture(str(caminho), width=LARGURA_UTIL)


def _remove(paragrafo) -> None:
    paragrafo._element.getparent().remove(paragrafo._element)


def _acha(doc: Document, teste) -> int | None:
    for i, p in enumerate(doc.paragraphs):
        if teste(p.text.strip()):
            return i
    return None


def monta(
    template: Path,
    markdown: Path,
    painel: Path,
    calendario: Path,
    destino: Path,
) -> Path:
    """Monta o .docx final e devolve o caminho gravado."""
    marcadores, ignoradas = le_marcadores(markdown)
    if not marcadores:
        raise RuntimeError(
            f"{markdown} não tem nenhum marcador. O comentário revisado é uma "
            "lista de marcadores iniciados por '- '."
        )

    if not (MIN_MARCADORES <= len(marcadores) <= MAX_MARCADORES):
        print(f"Aviso: o comentário tem {len(marcadores)} marcadores, fora da faixa "
              f"de {MIN_MARCADORES} a {MAX_MARCADORES} que o guia de estilo fixa. "
              "Excesso e falta apontam problema na redação ou na revisão, não "
              "caso legítimo de formato.", file=sys.stderr)

    fecho = [linha for linha in ignoradas if RE_FECHO.match(linha)]
    if fecho:
        print(f"Aviso: o comentário traz fecho próprio ({fecho[0]!r}), que foi "
              "descartado — o fecho vem do template, para não sair duplicado.",
              file=sys.stderr)
    resto = [linha for linha in ignoradas if not RE_FECHO.match(linha)]
    if resto:
        print(f"Aviso: {len(resto)} linha(s) fora de marcador foram ignoradas. "
              f"Primeira: {resto[0][:70]!r}", file=sys.stderr)

    doc = abre_template(template, destino)

    # As duas imagens, nas posições que o template nomeia. O painel ocupa o lugar
    # da tabela de fechamento, no topo; o calendário fica depois dos marcadores.
    i_painel = _acha(doc, lambda t: t == MARCA_PAINEL)
    if i_painel is None:
        raise RuntimeError(f"O template não tem o marcador {MARCA_PAINEL!r}.")
    _imagem(doc.paragraphs[i_painel], painel)

    i_calendario = _acha(doc, lambda t: t == MARCA_CALENDARIO)
    if i_calendario is None:
        raise RuntimeError(f"O template não tem o marcador {MARCA_CALENDARIO!r}.")
    _imagem(doc.paragraphs[i_calendario], calendario)

    # Parágrafos de conteúdo: um por marcador. Cada um leva consigo o espaçador
    # que o segue, para que a clonagem preserve o ritmo vertical do template.
    modelos = [i for i, p in enumerate(doc.paragraphs) if RE_PARAGRAFO.match(p.text.strip())]
    if not modelos:
        raise RuntimeError("O template não tem parágrafos de exemplo '[Parágrafo N – …]'.")

    alvos = [doc.paragraphs[i] for i in modelos]

    if len(marcadores) > len(alvos):
        ultimo = alvos[-1]
        espacador = ultimo._element.getnext()
        ancora = ultimo
        for _ in range(len(marcadores) - len(alvos)):
            if espacador is not None:
                novo_esp = copy.deepcopy(espacador)
                ancora._element.addnext(novo_esp)
                alvo_el = copy.deepcopy(ultimo._element)
                novo_esp.addnext(alvo_el)
            else:
                alvo_el = copy.deepcopy(ultimo._element)
                ancora._element.addnext(alvo_el)
            from docx.text.paragraph import Paragraph
            ancora = Paragraph(alvo_el, ultimo._parent)
            alvos.append(ancora)

    for alvo, trechos in zip(alvos, marcadores):
        _escreve(alvo, trechos)

    # Sobras: remove o parágrafo de exemplo e o espaçador que o acompanha.
    for alvo in alvos[len(marcadores):]:
        seguinte = alvo._element.getnext()
        if seguinte is not None and seguinte.tag.endswith("}p") and not "".join(
            seguinte.itertext()
        ).strip():
            seguinte.getparent().remove(seguinte)
        _remove(alvo)

    doc.save(str(destino))
    return destino
