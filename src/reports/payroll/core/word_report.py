"""
Word document generator for the Payroll report.

Generates a complete .docx report following the reference layout (5 pages)
using python-docx from scratch (no template).
"""

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from reports import _paths

from .calculations import MESES_PTBR_FULL

# === Document styling constants ===
FONT_NAME = "Calibri"
HEADER_RGB = RGBColor(0x2E, 0x4C, 0x59)  # COPOM[0]
ALT_ROW_HEX = "F5F0E8"
WHITE_HEX = "FFFFFF"
HEADER_HEX = "2E4C59"


def _apply_cell_shading(cell, hex_color: str) -> None:
    """Apply background shading to a table cell via XML."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def _setup_document() -> Document:
    """Create a new Document with proper margins and font defaults."""
    doc = Document()

    # Set margins (2cm top/bottom, 2cm left/right)
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = FONT_NAME
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    return doc


def _add_title_header(doc: Document, release_date: str) -> None:
    """Add the report title header."""
    dt = pd.to_datetime(release_date)
    month_name = MESES_PTBR_FULL.get(dt.month, "")
    year = dt.year

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"DADOS DO RELATORIO DE EMPREGO DOS EUA - {month_name.upper()}/{year}")
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = FONT_NAME
    run.font.color.rgb = HEADER_RGB

    # Add thin line
    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = line.add_run("_" * 80)
    run.font.size = Pt(6)
    run.font.color.rgb = RGBColor(0xAE, 0xAE, 0xAE)


def _add_release_table(doc: Document, df: pd.DataFrame) -> None:
    """Add a formatted release summary table to the document."""
    n_rows = len(df)
    n_cols = len(df.columns)

    table = doc.add_table(rows=n_rows + 1, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Header row
    for j, col_name in enumerate(df.columns):
        cell = table.rows[0].cells[j]
        cell.text = col_name
        _apply_cell_shading(cell, HEADER_HEX)
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.runs[0]
        run.bold = True
        run.font.size = Pt(10)
        run.font.name = FONT_NAME
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Data rows
    for i, (_, row) in enumerate(df.iterrows()):
        bg = ALT_ROW_HEX if i % 2 == 1 else WHITE_HEX
        for j, col in enumerate(df.columns):
            cell = table.rows[i + 1].cells[j]
            val = row[col]
            cell.text = "" if pd.isna(val) else str(val)
            _apply_cell_shading(cell, bg)

            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.size = Pt(10)
                run.font.name = FONT_NAME
                if col == "Atual":
                    run.bold = True


def _add_chart_page(doc: Document, image_path: Path, width_inches: float = 6.5) -> None:
    """Add a page break and insert a chart image."""
    doc.add_page_break()
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))


def _add_chart_inline(doc: Document, image_path: Path, width_inches: float = 6.5) -> None:
    """Insert a chart image without page break."""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))


def _add_headlines_section(doc: Document, headlines_text: str) -> None:
    """Add Bloomberg headlines section with formatted paragraphs."""
    doc.add_page_break()

    title = doc.add_paragraph()
    run = title.add_run("HEADLINES")
    run.bold = True
    run.font.size = Pt(13)
    run.font.name = FONT_NAME
    run.font.color.rgb = HEADER_RGB

    for line in headlines_text.strip().split("\n"):
        line = line.strip()
        if not line:
            continue
        para = doc.add_paragraph()
        # Lines starting with * or - are bold (key headlines)
        if line.startswith(("*", "-", ">")):
            line_clean = line.lstrip("*-> ").strip()
            run = para.add_run(line_clean)
            run.bold = True
        else:
            run = para.add_run(line)
        run.font.size = Pt(10)
        run.font.name = FONT_NAME


def generate_payroll_report(
    release_date: str,
    release_summary: pd.DataFrame,
    chart_paths: dict[str, Path],
    summary_text: str = "",
    headlines_text: str = "",
    market_reaction_path: Path | None = None,
    output_dir: Path | None = None,
) -> Path:
    """Generate the complete Payroll Word report.

    Args:
        release_date: Release date string (YYYYMMDD or YYYY-MM-DD).
        release_summary: DataFrame from format_release_summary().
        chart_paths: Dict mapping chart names to file paths. Expected keys:
            'nfp_chart', 'dashboard', 'industry'. Optional: 'u6_chart', 'beveridge'.
        summary_text: Opening summary paragraph (pt-BR).
        headlines_text: Bloomberg headlines text block.
        market_reaction_path: Optional PNG of market reaction (user-provided).
        output_dir: Directory for output. Defaults to output/reports/payroll/.

    Returns:
        Path to the generated .docx file.
    """
    if output_dir is None:
        output_dir = _paths.OUTPUT / "reports" / "payroll"
    output_dir.mkdir(parents=True, exist_ok=True)

    release_date_clean = release_date.replace("-", "")

    doc = _setup_document()

    # === Page 1: Title + Summary + Release Table ===
    _add_title_header(doc, release_date)

    # Summary text
    if summary_text:
        para = doc.add_paragraph()
        run = para.add_run(summary_text)
        run.font.size = Pt(11)
        run.font.name = FONT_NAME

    doc.add_paragraph()  # spacing

    # Release table
    _add_release_table(doc, release_summary)

    # === Page 2: Industry Breakdown ===
    if "industry" in chart_paths and chart_paths["industry"].exists():
        _add_chart_page(doc, chart_paths["industry"])

    # === Page 3: NFP Time Series + Dashboard ===
    if "nfp_chart" in chart_paths and chart_paths["nfp_chart"].exists():
        _add_chart_page(doc, chart_paths["nfp_chart"])

    if "dashboard" in chart_paths and chart_paths["dashboard"].exists():
        _add_chart_inline(doc, chart_paths["dashboard"])

    # === Extended charts (U6, Beveridge) ===
    has_extended = False
    for key in ("u6_chart", "beveridge"):
        if key in chart_paths and chart_paths[key].exists():
            if not has_extended:
                doc.add_page_break()
                title_para = doc.add_paragraph()
                run = title_para.add_run("INDICADORES COMPLEMENTARES")
                run.bold = True
                run.font.size = Pt(13)
                run.font.name = FONT_NAME
                run.font.color.rgb = HEADER_RGB
                has_extended = True
            _add_chart_inline(doc, chart_paths[key])

    # === Page 4: Market Reaction (optional) ===
    if market_reaction_path and market_reaction_path.exists():
        _add_chart_page(doc, market_reaction_path)

    # === Page 5: Headlines (optional) ===
    if headlines_text.strip():
        _add_headlines_section(doc, headlines_text)

    # Save
    filename = f"Mesa de Investimentos - PAYROLL {release_date_clean}.docx"
    output_path = output_dir / filename
    doc.save(str(output_path))

    return output_path
