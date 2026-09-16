"""Testes de caracterização de reports.fomc.

Fixam o comportamento atual antes de o pacote migrar para o repositório
mkt_intelligence; lá, os mesmos testes provam que a mudança de casa não mudou nada.
"""

from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
import pytest
from docx import Document

from reports.fomc.core import data_loader, word_export
from reports.fomc.core.calculations import calculate_surprise, classify_change_direction
from reports.fomc.core.data_loader import _TZ_BRT, _TZ_ET, is_sep_meeting
from reports.fomc.core.pdf_parser import (
    _parse_rate_fraction,
    _parse_statement_text,
    _parse_variable_tokens,
)
from reports.fomc.core.word_export import (
    MARKET_REACTION_PANELS,
    _to_naive_brt,
    create_market_reaction_grid,
)
from reports.fomc.core.word_report import _add_marked_text, _add_summary

STATEMENT = (
    "Inflation remains somewhat elevated. The Committee decided to maintain the target "
    "range for the federal funds rate at 4-1/4 to 4-1/2 percent. Voting against the "
    "action were: Michelle Bowman, Christopher Waller."
)


class TestClassifyChangeDirection:
    @pytest.mark.parametrize(
        ("variable", "change", "expected"),
        [
            ("PCE inflation", 0.2, "hawkish"),
            ("PCE inflation", -0.2, "dovish"),
            ("Unemployment rate", 0.1, "dovish"),
            ("Unemployment rate", -0.1, "hawkish"),
            ("Change in real GDP", 0.0005, "neutral"),
            ("Indicador desconhecido", 0.25, "hawkish"),
        ],
    )
    def test_direction(self, variable, change, expected):
        """Desemprego inverte o sinal; indicador fora do mapa sobe como hawkish."""
        assert classify_change_direction(variable, change) == expected


class TestCalculateSurprise:
    def test_rate_surprise_in_bps(self):
        """Surpresa de taxa sai também em pontos-base inteiros."""
        result = calculate_surprise(4.5, 4.25, "rate")
        assert result["surprise"] == 0.25
        assert result["surprise_bps"] == 25
        assert result["direction"] == "forte"

    def test_unemployment_is_inverted(self):
        """Desemprego abaixo do esperado é leitura forte, e não tem bps."""
        result = calculate_surprise(4.1, 4.3, "unemployment")
        assert abs(result["surprise"] - (-0.2)) < 1e-9
        assert result["surprise_bps"] is None
        assert result["direction"] == "forte"

    def test_no_surprise_is_neutral(self):
        """Atual igual à pesquisa é neutro."""
        assert calculate_surprise(2.0, 2.0, "inflation")["direction"] == "neutro"

    @pytest.mark.parametrize(
        ("actual", "survey", "expected"),
        [(4.26, 4.25, 1), (4.25, 4.26, -1)],
    )
    def test_one_bp_surprise(self, actual, survey, expected):
        """Surpresa de um ponto-base sai como um ponto-base, nos dois sentidos.

        `actual - survey` não dá 0,01 exato em ponto flutuante: dá 0,00999…, e
        multiplicado por 100 vira 0,9999…. Truncar isso apaga a surpresa. Apaga
        também a de −1bp, porque o truncamento anda em direção ao zero — e é por
        isso que o caso negativo está aqui.
        """
        assert calculate_surprise(actual, survey, "rate")["surprise_bps"] == expected


class TestIsSepMeeting:
    def test_sep_months(self):
        """Março, junho, setembro e dezembro têm SEP; janeiro não."""
        assert is_sep_meeting("20260318") is True
        assert is_sep_meeting("20260128") is False


class TestParseRateFraction:
    @pytest.mark.parametrize(("text", "expected"), [("4-1/4", 4.25), ("3-1/2", 3.5), ("4", 4.0)])
    def test_fraction(self, text, expected):
        """Frações do statement viram decimal."""
        assert _parse_rate_fraction(text) == expected


class TestParseVariableTokens:
    """Em setembro o SEP ganha um ano; a linha da projeção anterior tem um a menos."""

    def test_prior_line_with_one_year_fewer_keeps_longer_run_aligned(self):
        """12 tokens numa tabela de 4 anos + LR: três medianas, LR na 4ª, sem 2029."""
        tokens = [
            "1.1",
            "1.2",
            "1.3",
            "1.9",
            "1.0–1.2",
            "1.1–1.3",
            "1.2–1.4",
            "1.8–2.0",
            "0.9–1.3",
            "1.0–1.4",
            "1.1–1.5",
            "1.7–2.1",
        ]
        parsed = _parse_variable_tokens(tokens, n_years=4, has_longer_run=True)
        assert parsed["medians"] == {0: 1.1, 1: 1.2, 2: 1.3, "lr": 1.9}
        assert parsed["ct"] == {0: "1.0–1.2", 1: "1.1–1.3", 2: "1.2–1.4", "lr": "1.8–2.0"}
        assert parsed["range"] == {0: "0.9–1.3", 1: "1.0–1.4", 2: "1.1–1.5", "lr": "1.7–2.1"}

    def test_prior_line_without_longer_run(self):
        """Core PCE: 9 tokens numa tabela de 4 anos sem LR ficam nos três primeiros anos."""
        tokens = [
            "2.1",
            "2.2",
            "2.3",
            "2.0–2.2",
            "2.1–2.3",
            "2.2–2.4",
            "1.9–2.3",
            "2.0–2.4",
            "2.1–2.5",
        ]
        parsed = _parse_variable_tokens(tokens, n_years=4, has_longer_run=False)
        assert parsed["medians"] == {0: 2.1, 1: 2.2, 2: 2.3}
        assert parsed["ct"] == {0: "2.0–2.2", 1: "2.1–2.3", 2: "2.2–2.4"}

    def test_full_line_unchanged(self):
        """Linha completa de 3 anos + LR continua como antes."""
        tokens = [
            "1.1",
            "1.2",
            "1.3",
            "1.9",
            "a–b",
            "c–d",
            "e–f",
            "g–h",
            "i–j",
            "k–l",
            "m–n",
            "o–p",
        ]
        parsed = _parse_variable_tokens(tokens, n_years=3, has_longer_run=True)
        assert parsed["medians"] == {0: 1.1, 1: 1.2, 2: 1.3, "lr": 1.9}
        assert parsed["range"] == {0: "i–j", 1: "k–l", 2: "m–n", "lr": "o–p"}


class TestParseStatementText:
    def test_extracts_range_decision_and_dissent(self):
        """Faixa, decisão, dissidentes e frase-chave saem do texto do statement."""
        result = _parse_statement_text(STATEMENT)
        assert result["target_rate_low"] == 4.25
        assert result["target_rate_high"] == 4.5
        assert result["decision"] == "hold"
        assert result["unanimous"] is False
        assert result["dissenters"] == ["Michelle Bowman", "Christopher Waller"]
        assert result["key_phrases"] == ["Inflation remains somewhat elevated"]


class TestToNaiveBrt:
    def test_winter_and_summer_offsets(self):
        """ET→BRT: +2h no inverno de Nova York, +1h no verão; sai sem fuso."""
        winter = _to_naive_brt(pd.Timestamp("2026-01-28 14:00", tz=_TZ_ET))
        summer = _to_naive_brt(pd.Timestamp("2026-07-29 14:00", tz=_TZ_ET))
        assert winter == pd.Timestamp("2026-01-28 16:00")
        assert summer == pd.Timestamp("2026-07-29 15:00")

    def test_naive_passes_through(self):
        """Timestamp sem fuso já é BRT e passa intacto."""
        ts = pd.Timestamp("2026-01-28 16:00")
        assert _to_naive_brt(ts) == ts


def _synthetic_intraday() -> pd.DataFrame:
    index = pd.date_range("2026-01-28 09:00", "2026-01-28 20:00", freq="5min", tz=_TZ_BRT)
    columns = [key for key, *_ in MARKET_REACTION_PANELS]
    return pd.DataFrame(
        {col: [100.0 + i * 0.01 for i in range(len(index))] for col in columns},
        index=index,
    )


class TestCreateMarketReactionGrid:
    @pytest.mark.parametrize("style", ["word", "bloomberg"])
    def test_writes_png(self, tmp_path, style):
        """O grid 3x3 sai em PNG nos dois estilos, com dados sintéticos."""
        out = tmp_path / f"grid_{style}.png"
        fig = create_market_reaction_grid(
            _synthetic_intraday(),
            event_times={"decisao": pd.Timestamp("2026-01-28 16:00", tz=_TZ_BRT)},
            output_path=out,
            style=style,
        )
        n_axes = len(fig.axes)
        plt.close(fig)
        assert out.stat().st_size > 0
        assert n_axes >= 9

    def test_rejects_unknown_style(self):
        """Estilo fora de MARKET_REACTION_STYLES é erro."""
        with pytest.raises(ValueError):
            create_market_reaction_grid(_synthetic_intraday(), style="excel")


class TestProjectRootAnchors:
    """input/, output/ e etc/ se ancoram na raiz do projeto — a que tem o pyproject.toml.

    A migração só é neutra se a raiz do produto novo repetir esse arranjo.
    """

    def test_word_export_root_has_pyproject(self):
        """output/reports/fomc nasce na raiz do projeto."""
        assert (word_export._get_project_root() / "pyproject.toml").is_file()

    def test_grid1_anchor_is_project_root(self):
        """O grid1.xlsx é procurado em <raiz>/input/."""
        assert (Path(data_loader.__file__).parents[4] / "pyproject.toml").is_file()

    def test_module_path_is_fomc_package(self):
        """Os documentos do Fed moram em <pacote fomc>/input/."""
        assert data_loader._get_module_path().name == "fomc"


class TestMarkedText:
    def test_italic_bold_and_plain_runs(self):
        """`*x*` vira itálico, `**y**` negrito, o resto fica normal."""
        doc = Document()
        para = doc.add_paragraph()
        _add_marked_text(para, "leitura *hawkish* e **firme** hoje")
        runs = [(r.text, bool(r.italic), bool(r.bold)) for r in para.runs]
        assert runs == [
            ("leitura ", False, False),
            ("hawkish", True, False),
            (" e ", False, False),
            ("firme", False, True),
            (" hoje", False, False),
        ]

    def test_summary_paragraphs_keep_italics(self):
        """O resumo separa parágrafos por linha em branco e aplica a marcação em cada um."""
        doc = Document()
        _add_summary(doc, "Primeiro *dots*.\n\nSegundo.")
        paras = doc.paragraphs
        assert [p.text for p in paras] == ["Primeiro dots.", "Segundo."]
        assert any(r.italic for r in paras[0].runs)
